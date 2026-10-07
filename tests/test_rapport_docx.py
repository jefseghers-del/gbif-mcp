"""Tests voor de Word-variant van het datarapport (gbif_mcp/rapport_docx.py) en `formaat` in de tool. Geen netwerk."""
from __future__ import annotations

import asyncio

import pytest
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt

from gbif_mcp import DISCLAIMER_KORT, PRIVACY, server
from gbif_mcp.rapport_docx import schrijf_docx
from tests.test_rapport import _data


def _alle_tekst(doc) -> str:
    delen = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for rij in t.rows:
            delen += [c.text for c in rij.cells]
    delen += [p.text for p in doc.sections[0].footer.paragraphs]
    return " ".join(" ".join(delen).split())


def _volledig() -> dict:
    """Invoer met BWK-tabel, dekking per soortgroep en onderliggende waarnemingen."""
    D = _data(straal_gebieden=200)
    D["gebieden"]["lagen"].append({
        "laag": "bwk_habitat", "naam": "BWK — habitat", "status": "ok", "aantal_binnen_straal": 1,
        "samenvatting_bwk": {"binnen_straal": {"totaal": 1, "waardevol": 0, "zeer_waardevol": 1},
                             "locatie_zelf": [{"label": "hf", "waardering": "biologisch zeer waardevol", "habitat": "rbbhf 100 %"}]},
        "treffers": [{"naam": "hf", "overlapt": True, "afstand_m": 0, "bwk": {
            "bwklabel": "hf", "eval": "z", "waardering": "biologisch zeer waardevol", "karteerjaar_of_versie": "v2014",
            "eenheden": [{"code": "hf", "omschrijving": "moerasspirearuigte"}], "kaartsleutel": "rbbhf",
            "habitats": [], "rbb": [{"code": "rbbhf", "aandeel_pct": 100}]}}],
    })
    D["gebieden"]["samenvatting"]["bwk"] = {
        "binnen_straal": {"totaal": 1, "waardevol": 0, "zeer_waardevol": 1},
        "locatie_zelf": [{"label": "hf", "waardering": "biologisch zeer waardevol", "habitat": "rbbhf 100 %"}],
        "habitattypes_binnen_straal": [], "rbb_binnen_straal": [], "dichtste_waardevol": None}
    D["kern"]["dekking"] = [{"groep": "vleermuizen", "wetenschappelijke_naam": "Chiroptera", "status": "ok", "totaal": 20,
                             "per_dataset": [{"dataset_key": "280674cb-42f8-4959-b6aa-eee663157965", "aantal": 17, "aandeel": 0.85}],
                             "waarschuwing": "Vleermuizen: 85% van de 20 GBIF-records in het gebied komt uit één INBO-dataset."}]
    D["detail"] = [{"soort": D["kern"]["soorten"][0], "dataset_key": "280674cb-42f8-4959-b6aa-eee663157965", "waarnemingen": {
        "totaal": 1, "per_verificatiestatus": {"(leeg)": 1}, "per_methode": {"batdetector": 1},
        "waarnemingen": [{"datum": "2025-06-01", "plaats": "Testplaats", "gemeente": None, "onzekerheid_m": 707.0,
                          "verificatiestatus": None, "basis": "HUMAN_OBSERVATION", "methode": "batdetector"}]}}]
    return D


@pytest.fixture
def docx_pad(tmp_path, monkeypatch):
    from gbif_mcp import bwk

    # habitat_tekst/wordt_getekend hangen af van de BWK-legende; hier volstaat een vaste weergave.
    monkeypatch.setattr(bwk, "habitat_tekst", lambda e: "rbbhf 100 %")
    monkeypatch.setattr(bwk, "wordt_getekend", lambda e, overlapt: True)
    pad = tmp_path / "r.docx"
    schrijf_docx(_volledig(), str(pad))
    return pad


def test_docx_bevat_hoofdstukken_bwk_dekking_en_disclaimer(docx_pad):
    doc = Document(str(docx_pad))
    tekst = _alle_tekst(doc)
    koppen = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for h in ("1. Samenvatting", "3. Statussen in cijfers", "4. Kernsoorten", "5. Beschermde gebieden en gebiedsstatuten",
              "5.1 Biologische Waarderingskaart: alle eenheden binnen de straal",
              "6. Onderliggende waarnemingen van de striktst beschermde soorten", "7. Verantwoording van de bronnen",
              "7.2bis Dekking per soortgroep", "8. Beperkingen"):
        assert h in koppen
    assert DISCLAIMER_KORT in tekst
    assert " ".join(PRIVACY.split()) in tekst
    assert "De projectlocatie ligt in BWK-eenheid hf" in tekst  # rapport.bwk_zin
    assert "moerasspirearuigte" in tekst and "v2014" in tekst  # tabel 5.1
    assert "85%" in tekst and "Vleermuizen: 85% van de 20" in tekst
    assert "batdetector" in tekst
    assert "12.345" in tekst and "Testwaarschuwing." in tekst
    assert "<b>" not in tekst and "<br" not in tekst and "<i>" not in tekst


def test_docx_markup_wordt_opmaak(docx_pad):
    doc = Document(str(docx_pad))
    vet = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "12.345" in vet  # <b>…</b> in de samenvatting


def test_docx_huisstijl_corpus(docx_pad):
    doc = Document(str(docx_pad))
    normaal = doc.styles["Normal"]
    assert normaal.font.name == "Calibri" and normaal.font.size == Pt(10)
    pf = normaal.paragraph_format
    assert pf.line_spacing_rule == WD_LINE_SPACING.EXACTLY and pf.line_spacing == Pt(15)
    assert pf.space_before == Pt(0) and pf.space_after == Pt(0)
    corpus = [p for p in doc.paragraphs if p.style.name == "Normal" and p.text.strip()]
    assert len(corpus) > 5
    for p in corpus:
        # Geen afwijkende opmaak op alineaniveau: de stijl bepaalt het.
        assert p.paragraph_format.line_spacing in (None, Pt(15))
        assert p.paragraph_format.space_before in (None, Pt(0)) and p.paragraph_format.space_after in (None, Pt(0))
        for r in p.runs:
            assert r.font.name in (None, "Calibri") and r.font.size in (None, Pt(10))
    for naam in ("Heading 1", "Heading 2", "Title", "Noot", "Cel"):
        assert doc.styles[naam].font.name == "Calibri"


def test_docx_bevat_geen_waarnemersnamen(tmp_path):
    D = _volledig()
    # Mocht een bron toch persoonsvelden meesturen: het rapport neemt ze nooit over.
    D["detail"][0]["waarnemingen"]["waarnemingen"][0]["recordedBy"] = "Jan Waarnemer"
    D["detail"][0]["waarnemingen"]["waarnemingen"][0]["identifiedBy"] = "Piet Determinator"
    pad = tmp_path / "r.docx"
    schrijf_docx(D, str(pad))
    tekst = _alle_tekst(Document(str(pad)))
    assert "Jan Waarnemer" not in tekst and "Piet Determinator" not in tekst


# Tool: parameter formaat --------------------------------------------------------------------------------


class _M:
    """Nabootsing van een Pydantic-model: attributen en model_dump."""

    def __init__(self, d: dict):
        self._d = d

    def __getattr__(self, k):
        v = self._d.get(k)
        return [_M(x) if isinstance(x, dict) else x for x in v] if isinstance(v, list) else v

    def model_dump(self, mode="json"):
        return self._d


def _nep_tool(monkeypatch):
    D = _data()

    async def telling(**kw):
        return _M(D["telling"])

    async def soorten(**kw):
        return _M({**D["kern"], "licentiefilter": "alle licenties", "uitgesloten_niet_commercieel": 0})

    async def gebieden(**kw):
        return _M({**D["gebieden"], "niet_geraadpleegd": []})

    monkeypatch.setattr(server, "telling_in_gebied", telling)
    monkeypatch.setattr(server, "soorten_in_gebied", soorten)
    monkeypatch.setattr(server, "gebieden_rond", gebieden)


@pytest.mark.parametrize("formaat,pdf,docx", [("pdf", True, False), ("docx", False, True), ("beide", True, True)])
def test_tool_formaat(tmp_path, monkeypatch, formaat, pdf, docx):
    _nep_tool(monkeypatch)
    r = asyncio.run(server.datarapport_natuur(lat=51.05, lon=3.72, pad=str(tmp_path / "rapport.pdf"), kaarten=False,
                                              detail_soorten=0, formaat=formaat))
    assert (r["pad"] is not None) == pdf and (tmp_path / "rapport.pdf").exists() == pdf
    assert (r["pad_docx"] is not None) == docx and (tmp_path / "rapport.docx").exists() == docx


def test_tool_onbekend_formaat(monkeypatch):
    _nep_tool(monkeypatch)
    with pytest.raises(ValueError, match="pdf, docx of beide"):
        asyncio.run(server.datarapport_natuur(lat=51.05, lon=3.72, formaat="odt"))
