# Copyright (c) 2026 Jef Seghers
# In licentie gegeven krachtens de EUPL
# SPDX-License-Identifier: EUPL-1.2
"""Datarapport natuur als Word-document (.docx), naast de PDF uit `rapport.py`.

Zelfde invoer (`D`), zelfde opbouw en dezelfde teksten: alle inhoud komt uit de gedeelde functies
in `rapport.py`. Deze module zorgt alleen voor de Word-opmaak in de LDR-huisstijl: corpustekst in
Calibri 10 pt, regelafstand exact 15 pt, 0 pt vóór en na de alinea. Koppen, noten en tabellen
sober, ook in Calibri. Geen logo of sjabloon: die zijn niet beschikbaar en worden niet nagebootst.

De beperkte markup van de gedeelde teksten (<b>, <i>, <br/>, <font …>) wordt omgezet in Word-runs;
<font> wordt genegeerd.
"""
from __future__ import annotations

import html
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from . import rapport as R

LETTERTYPE = "Calibri"
CORPUS_PT = 10
REGELAFSTAND_PT = 15
GRIJS = RGBColor(0x4A, 0x4A, 0x4A)
KOPGRIJS = "D9D9D9"
LIJN = "A6A6A6"


# ---------------------------------------------------------------- stijlen


def _lettertype(stijl, grootte: float, vet: bool = False, kleur: RGBColor | None = None) -> None:
    f = stijl.font
    f.name = LETTERTYPE
    f.size = Pt(grootte)
    f.bold = vet
    f.italic = False
    if kleur is not None:
        f.color.rgb = kleur
    # Ook voor Oost-Aziatische en complexe tekens Calibri, anders valt Word terug op het themalettertype.
    rpr = stijl.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), LETTERTYPE)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)


def _alinea(stijl, regel_pt: float, voor: float = 0, na: float = 0) -> None:
    pf = stijl.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(regel_pt)
    pf.space_before = Pt(voor)
    pf.space_after = Pt(na)


def _stijlen(doc) -> None:
    st = doc.styles
    normaal = st["Normal"]
    _lettertype(normaal, CORPUS_PT)
    _alinea(normaal, REGELAFSTAND_PT)

    for naam, grootte, voor, na, kleur in (("Title", 18, 0, 0, None), ("Heading 1", 12, 15, 0, None), ("Heading 2", 10, 15, 0, GRIJS)):
        s = st[naam]
        _lettertype(s, grootte, vet=True, kleur=kleur or RGBColor(0, 0, 0))
        _alinea(s, max(REGELAFSTAND_PT, grootte + 4), voor, na)
        s.paragraph_format.keep_with_next = True
        # Titelstijl van het standaardsjabloon heeft een onderrand; weg ermee (sober).
        ppr = s.element.get_or_add_pPr()
        for rand in ppr.findall(qn("w:pBdr")):
            ppr.remove(rand)

    for naam, grootte, regel in (("Noot", 8, 11), ("Cel", 8, 10)):
        if naam not in [x.name for x in st]:
            s = st.add_style(naam, 1)  # 1 = alineastijl
            s.base_style = normaal
        s = st[naam]
        _lettertype(s, grootte, kleur=GRIJS if naam == "Noot" else None)
        _alinea(s, regel)


# ---------------------------------------------------------------- markup -> runs

_TAG = re.compile(r"(<br\s*/?>|</?b>|</?i>|<font[^>]*>|</font>)", re.I)


def _runs(par, tekst: str, vet: bool = False) -> None:
    """Zet de beperkte markup van rapport.py om in runs van `par`."""
    b, i = vet, False
    for deel in _TAG.split(str(tekst)):
        if not deel:
            continue
        t = deel.lower()
        if t.startswith("<br"):
            par.add_run().add_break()
        elif t == "<b>":
            b = True
        elif t == "</b>":
            b = vet
        elif t == "<i>":
            i = True
        elif t == "</i>":
            i = False
        elif t.startswith("<font") or t == "</font>":
            continue
        else:
            r = par.add_run(html.unescape(deel))
            r.bold = b or None
            r.italic = i or None


def _p(doc, tekst: str = "", stijl: str = "Normal"):
    par = doc.add_paragraph(style=stijl)
    if tekst:
        _runs(par, tekst)
    return par


def _wit(doc) -> None:
    """Lege corpusregel als scheiding (de corpustekst zelf heeft 0 pt spatiëring)."""
    doc.add_paragraph(style="Normal")


def _pagina(doc) -> None:
    doc.add_paragraph(style="Normal").add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- tabellen


def _arcering(cel, kleur: str) -> None:
    tcpr = cel._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), kleur.lstrip("#").upper())
    tcpr.append(shd)


def _randen(tabel, kleur: str = LIJN, binnen: bool = True) -> None:
    tblpr = tabel._tbl.tblPr
    randen = OxmlElement("w:tblBorders")
    for kant in ("top", "left", "bottom", "right") + (("insideH", "insideV") if binnen else ()):
        e = OxmlElement(f"w:{kant}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), kleur)
        randen.append(e)
    tblpr.append(randen)


def _celmarges(tabel, mm: float = 1.0) -> None:
    """Smallere celmarges dan de Word-standaard (1,9 mm), zodat smalle cijferkolommen niet afbreken."""
    mar = OxmlElement("w:tblCellMar")
    for kant in ("left", "right"):
        e = OxmlElement(f"w:{kant}")
        e.set(qn("w:w"), str(round(mm * 56.7)))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tabel._tbl.tblPr.append(mar)


def _koprij_herhalen(rij) -> None:
    trpr = rij._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trpr.append(e)


def _celtekst(cel, tekst: str, vet: bool = False) -> None:
    par = cel.paragraphs[0]
    par.style = "Cel"
    _runs(par, tekst, vet=vet)


def _breedtes(tabel, breedtes_pt: list[float]) -> None:
    """Kolombreedtes in dezelfde verhouding als in de PDF, over de volle tekstbreedte (170 mm)."""
    totaal = sum(breedtes_pt)
    for i, b in enumerate(breedtes_pt):
        w = Mm(170 * b / totaal)
        tabel.columns[i].width = w
        for cel in tabel.columns[i].cells:
            cel.width = w


def _tabel(doc, kop: list[str], rijen: list[list[str]], breedtes: list[float]):
    t = doc.add_table(rows=1, cols=len(kop))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    _randen(t)
    _celmarges(t)
    for cel, k in zip(t.rows[0].cells, kop):
        _celtekst(cel, k, vet=True)
        _arcering(cel, KOPGRIJS)
    _koprij_herhalen(t.rows[0])
    for r in rijen:
        cellen = t.add_row().cells
        for cel, c in zip(cellen, r):
            _celtekst(cel, str(c))
    _breedtes(t, breedtes)
    return t


def _kv(doc, rijen: list[tuple[str, str]], breedte: float = 168):
    t = doc.add_table(rows=0, cols=2)
    t.autofit = False
    _celmarges(t, 0)
    for k, v in rijen:
        cellen = t.add_row().cells
        _celtekst(cellen[0], k, vet=True)
        _celtekst(cellen[1], v)
    _breedtes(t, [breedte, 460 - breedte])
    return t


def _kader(doc, tekst: str):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    _randen(t, kleur="E69F00", binnen=False)
    cel = t.rows[0].cells[0]
    _arcering(cel, "FFF4E0")
    _celtekst(cel, tekst)
    _breedtes(t, [1])
    return t


def _legende(doc, regels: list[dict]):
    """Kaartlegende: nummer in een gekleurd vakje (zelfde kleur als op de kaart) en de laagnaam.
    Het nummer staat ook in elk vlak op de kaart, zodat de legende zonder kleur leesbaar blijft."""
    t = doc.add_table(rows=0, cols=2)
    t.autofit = False
    for r in regels:
        cellen = t.add_row().cells
        if r.get("nummer"):
            _celtekst(cellen[0], str(r["nummer"]), vet=True)
            cellen[0].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _arcering(cellen[0], r["kleur"])
        naam = R.legendenaam(r)
        _celtekst(cellen[1], (f"<b>{r['nummer']}.</b> " if r.get("nummer") else "") + naam)
    _breedtes(t, [12, 448])
    return t


# ---------------------------------------------------------------- voettekst


def _veld(par, instructie: str) -> None:
    for soort, tekst in (("begin", None), (None, instructie), ("end", None)):
        run = par.add_run()
        if soort:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), soort)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = f" {tekst} "
        run._r.append(e)


def _voet(doc, D: dict) -> None:
    sectie = doc.sections[0]
    voet = sectie.footer
    par = voet.paragraphs[0]
    par.style = "Noot"
    par.add_run(R.voettekst(D) + " — p. ")
    _veld(par, "PAGE")
    par.add_run().add_break()
    par.add_run(R.VOETTEKST_2)


# ---------------------------------------------------------------- document


def schrijf_docx(D: dict, pad: str) -> dict:
    """Schrijf het datarapport natuur als .docx naar `pad`. Zelfde invoer als `rapport.schrijf_pdf`."""
    doc = Document()
    sectie = doc.sections[0]
    sectie.orientation = WD_ORIENT.PORTRAIT
    sectie.page_width, sectie.page_height = Mm(210), Mm(297)
    sectie.left_margin = sectie.right_margin = Mm(20)
    sectie.top_margin, sectie.bottom_margin = Mm(18), Mm(20)
    _stijlen(doc)
    doc.core_properties.title = "Datarapport natuur"
    doc.core_properties.author = "MCP-connector BE-biodiversiteit"

    tel, kern, geb = D["telling"], D["kern"], D["gebieden"]
    _, straal_gebieden = R.stralen(D)

    # ---------------------------------------------------------------- titelblad + samenvatting
    _p(doc, "Datarapport natuur", "Title")
    ondertitel = _p(doc, R.ondertitel(D))
    for r in ondertitel.runs:
        r.font.color.rgb = GRIJS
    _wit(doc)
    _kv(doc, R.titelgegevens(D))
    _wit(doc)
    _kader(doc, R.KADER)
    _p(doc, "1. Samenvatting", "Heading 1")
    for i, t in enumerate(R.samenvatting(D)):
        if i:
            _wit(doc)
        _p(doc, t)

    # ---------------------------------------------------------------- 2. situering
    kaarten = R.kaartlijst(D)
    if kaarten:
        _pagina(doc)
        _p(doc, "2. Situering", "Heading 1")
        _p(doc, R.situering(kaarten))
        for i, k in enumerate(kaarten, start=1):
            _p(doc, f"Figuur {i} — {k.get('titel') or f'Kaart {i}'}", "Heading 2")
            beeld = doc.add_paragraph(style="Normal")
            # De afbeelding is hoger dan één regel: exacte regelafstand zou haar afknippen.
            beeld.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            beeld.add_run().add_picture(k["pad"], width=Mm(162 if len(kaarten) == 1 else 134))
            _legende(doc, k["legende"])
        _wit(doc)
        _p(doc, kaarten[0]["kanttekening"], "Noot")

    # ---------------------------------------------------------------- 3. statussen
    _pagina(doc)
    _p(doc, "3. Statussen in cijfers", "Heading 1")
    _tabel(doc, ["Lijst", "Categorie", "Aantal soorten"], R.statussen_rijen(tel), [170, 200, 90])
    _p(doc, R.STATUSSEN_NOOT, "Noot")

    # ---------------------------------------------------------------- 4. kernsoorten
    _pagina(doc)
    _p(doc, "4. Kernsoorten", "Heading 1")
    _p(doc, R.kernsoorten_inleiding(D))
    _wit(doc)
    _tabel(doc, R.KERNSOORTEN_KOP, R.kernsoorten_rijen(kern), [100, 110, 18, 28, 26, 24, 154])
    _p(doc, R.kernsoorten_noot(kern), "Noot")

    # ---------------------------------------------------------------- 5. gebieden
    _pagina(doc)
    rijen, bwk_laag = R.gebieden_rijen(geb)
    _p(doc, "5. Beschermde gebieden en gebiedsstatuten", "Heading 1")
    _p(doc, R.gebieden_inleiding(geb))
    _wit(doc)
    _tabel(doc, ["Laag", "Resultaat"], rijen, [175, 285])
    _p(doc, geb["kanttekening"], "Noot")
    if bwk_laag and bwk_laag["treffers"]:
        _p(doc, R.BWK_TITEL, "Heading 2")
        _p(doc, R.bwk_inleiding(bwk_laag, straal_gebieden), "Noot")
        _tabel(doc, R.BWK_KOP, R.bwk_rijen(bwk_laag, kaarten), [30, 38, 56, 156, 72, 70, 38])
        for n in R.bwk_noten(bwk_laag):
            _p(doc, n, "Noot")
    eco_laag = R.ecotoop_laag(geb)
    if eco_laag:
        _p(doc, R.ECOTOOP_TITEL, "Heading 2")
        _p(doc, R.ecotoop_inleiding(eco_laag, straal_gebieden), "Noot")
        _tabel(doc, R.ECOTOOP_KOP, R.ecotoop_rijen(eco_laag), [40, 70, 92, 72, 72, 72, 42])
        for n in R.ecotoop_noten(eco_laag):
            _p(doc, n, "Noot")

    # ---------------------------------------------------------------- 6. onderliggende waarnemingen
    _p(doc, R.DETAIL_TITEL, "Heading 1")
    _p(doc, R.DETAIL_INLEIDING)
    for blok in D["detail"]:
        titel, noot, drijen = R.detail_blok(blok)
        _p(doc, titel, "Heading 2")
        _p(doc, noot, "Noot")
        _tabel(doc, R.DETAIL_KOP, drijen, [52, 158, 38, 100, 112])

    # ---------------------------------------------------------------- 7. verantwoording
    _pagina(doc)
    _p(doc, "7. Verantwoording van de bronnen", "Heading 1")
    _p(doc, "7.1 Geraadpleegde soortenlijsten", "Heading 2")
    _tabel(doc, ["Code", "Lijst", "Opgehaald op"], R.lijst_rijen(kern), [80, 270, 110])
    _p(doc, R.rodelijst_noot(kern), "Noot")
    _p(doc, "7.2 Brondatasets in het zoekgebied", "Heading 2")
    _p(doc, R.datasets_inleiding(D), "Noot")
    _tabel(doc, ["Dataset", "Licentie", "Records", "Soorten"], R.dataset_rijen(tel), [270, 70, 60, 60])
    _p(doc, R.licentie_noot(kern), "Noot")
    if kern.get("dekking"):
        _p(doc, "7.2bis Dekking per soortgroep", "Heading 2")
        _p(doc, R.dekking_inleiding(D), "Noot")
        _tabel(doc, R.DEKKING_KOP, R.dekking_rijen(kern), [70, 45, 150, 195])
    _p(doc, "7.3 Kaartlagen", "Heading 2")
    lagen = R.kaartlagen(D)
    if lagen:
        _tabel(doc, lagen[0], lagen[1], [260, 90, 110])
    else:
        _p(doc, "Geen kaart opgenomen.", "Noot")
    _p(doc, "7.4 Reproduceerbaarheid", "Heading 2")
    _kv(doc, R.reproduceerbaarheid(kern), 120)

    # ---------------------------------------------------------------- 8. beperkingen
    _p(doc, "8. Beperkingen", "Heading 1")
    for i, zin in enumerate(R.beperkingen(kern)):
        if i:
            _wit(doc)
        _p(doc, "• " + zin)

    _voet(doc, D)
    doc.save(pad)
    return {"pad": pad}
